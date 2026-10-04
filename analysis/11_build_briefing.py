"""Builds the team briefing / judge-prep guide (out/Team_Briefing_Judge_Prep.docx).

Plain-language explanation of the project, the night's problems and fixes, the accuracy results,
and judge preparation. Numbers come from out/final_metrics.json. A tiny markdown-like syntax is
used below: '# ', '## ', '### ' headings, '- ' bullets, '> ' callouts, '|' table rows,
'**bold**' inline, blank line = new paragraph.
"""
import json
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from common import OUT

M = json.load(open(os.path.join(OUT, "final_metrics.json"), encoding="utf-8"))
dr, vr, mcd, mcv = M["distance_real"], M["velocity_real"], M["mc_distance"], M["mc_velocity"]
ins, loo, env = dr["final_in_sample"], dr["final_leave_one_out"], vr["envelope_stats"]
S2, VM, seg = mcd["new_run"], mcv["main"], vr["segments"]
pre, first = dr["before_any_fix_pct"], dr["after_first_fix_pct"]
st = M["story"]
VDRAFT = "PROVISIONAL" in json.dumps(mcv) or not mcv["sensitivity"]


def p(x, d=1):
    return f"{x:.{d}f}".replace("-", "−")


TEXT = f"""
# Open Ground Speed Sensor: team briefing and judge prep

> How to read this: Part 1 explains the project from zero. Part 2 is the story of the night (what broke, why, how it was fixed). Part 3 is how accurate it really is. Part 4 is judge preparation: the pitch, the demo, the hard questions with answers, and the claims to avoid. Part 5 is a one-page cheat sheet and a glossary. Part 6 is the technical deep dive: sensor fusion, each sensor's limits, the Kalman filter and the Monte Carlo, with every parameter and why it was chosen. Every number here comes from the re-analysis of the logged data in the repo (branch analysis/forensics-monte-carlo, file analysis/out/final_metrics.json).{" The velocity Monte Carlo numbers are from a provisional run (5,000 of 100,000 trials); the full run gave the same main result." if VDRAFT else ""}

# Part 1: What we built, from zero

## The 30-second version

We turned an iPhone into an optical ground-speed sensor. Point the phone's camera at the ground, light it with the flashlight, and the phone measures how fast the ground slides past, the same idea as an optical computer mouse. LiDAR measures how high the camera is above the ground, which converts "pixels per frame" into "metres per second". A Kalman filter on the phone blends that with the motion sensors and GPS, then streams speed and distance live to a laptop dashboard.

Professional optical ground-speed sensors (brands such as Kistler Correvit and Datron) cost thousands of dollars and are standard equipment for vehicle testing and Formula SAE. They measure true speed over the ground, which wheel speed sensors can't when tyres slip. This project does the core job with a phone people already own.

## Why measuring speed is harder than it sounds

- **Wheel speed sensors** count wheel rotations. If a tyre slips (hard acceleration, braking, cornering) the wheel speed is not the car's speed. Measuring that slip is the whole point of a ground-speed sensor.
- **GPS** gives speed, but only about once a second, with a lag, and phone GPS speed is noisy (in our data it trailed the camera by about 0.65 s).
- **Accelerometers** (the IMU) measure acceleration. You can add it up to get speed, but tiny errors accumulate, so after a few seconds the answer drifts away. We saw a phone sitting still "travel" 15 m in 3 s this way.
- **Looking at the ground** measures the actual motion directly, many times a second. That's why commercial sensors do it, and why we did.

## How each piece works (no maths needed)

### 1. The camera measures how far the ground moved between frames

The phone films the ground at 240 frames per second with a very short exposure (1/1000 s, to avoid blur) and the torch on. For every frame it takes the small square in the middle of the image and asks: "how far did this pattern shift since the last frame?"

It answers with **phase correlation**, a technique based on the Fourier transform. You don't need the maths: it compares the two images all at once and produces a sharp peak at the shift that lines them up best. Two things come out:
- the **shift** in pixels (with sub-pixel precision), and
- a **quality score called PSR** (peak-to-sidelobe ratio): how much the winning peak stands out from the noise. Around 3–7 means "I see nothing useful" (dark, blurry, blank floor); tens to hundreds means "clear, confident match".

Frames with PSR below 8 are thrown away. Frames with higher PSR are trusted more.

### 2. Pixels become metres using the camera height

A shift of 10 pixels means very different speeds depending on how far away the ground is: from high up, 10 pixels covers more ground. So we need two numbers:
- **h**, the camera's height above the ground, measured by the iPhone's LiDAR sensor before each run ("Measure h" button), and
- **f**, the camera's focal length in pixels (how zoomed-in the lens is).

The rule is: **speed = shift × (h ÷ f) ÷ time between frames**. Any error in h or f multiplies every speed and every distance. A 1 mm error in height at a 19 cm mount is a 0.5 % error in everything. This turned out to be the most important fact of the whole project.

Two small corrections also happen here. If the phone is tilted, the camera sees the ground at an angle, so the gravity sensor is used to correct for the tilt. If the phone rotates (wobbles on its mount), the image sweeps across the ground even though the car isn't moving, so the gyroscope is used to subtract that rotation ("de-rotation").

### 3. The Kalman filter blends everything into one best estimate

A **Kalman filter** is an algorithm that keeps a running best guess of something (here, speed) along with an honest estimate of how uncertain that guess is. It repeats two steps:
- **Predict** (100 times a second): use the accelerometer to guess how the speed changed since last time. Uncertainty grows a little, because accelerometers drift.
- **Correct** (every time a measurement arrives): compare the prediction to the camera's measured speed and move the estimate toward it. How far it moves depends on which is more trustworthy right now: a sharp camera frame (high PSR) pulls hard, a weak one barely pulls.

Our filter tracks four things (its **state**): forward speed, sideways speed, and the accelerometer's bias (its constant error) in each direction. Estimating the bias lets the filter learn and cancel the accelerometer's error over time.

Extra rules the filter uses:
- **Outlier rejection (gating):** if a camera reading is wildly inconsistent with the prediction (more than 3 standard deviations), it's ignored. If 12 readings in a row get rejected, the filter assumes it is the filter that is wrong and accepts the camera again.
- **Zero-velocity updates (ZUPT):** when the motion sensors say the phone is perfectly still, the filter is told "speed is exactly zero". This stops drift whenever the vehicle stops.
- **GPS speed** is also fed in when moving faster than 1 m/s. In practice it barely matters: the camera is trusted far more (removing GPS changes the estimate by at most 4 mm/s).

### 4. Distance is speed added up over time

Distance is the running total of speed × time. Two refinements were added after real tests: speeds below 0.05 m/s count as zero (so sensor noise at rest doesn't creep up), and speed is smoothed over half a second before adding (so a wobbling phone doesn't count every wobble as distance). A small correction term removes the half-second delay that the smoothing would otherwise add.

### 5. The dashboard

The phone sends 50 updates a second over Wi-Fi to a Python dashboard on a laptop: live speed, distance, quality, camera height, buttons for Measure h / Zero / Start / Stop, and a table of runs. Every run is also recorded on the phone at full rate (every camera frame, every IMU sample, every GPS fix, every filter output). That recording turned out to be the team's best decision: every bug in Part 2 was found in those logs.

## The system at a glance

| Piece | What it does | Rate |
| Camera + phase correlation | ground shift + quality (PSR) | 240 per second |
| LiDAR | camera height h, measured once before a run | once |
| IMU (accelerometer + gyroscope) | predicts speed changes; detects stillness; de-rotation | 100 per second |
| GPS | speed, used only above 1 m/s | 1 per second |
| Kalman filter (4 states) | best estimate of speed + uncertainty | 100 per second |
| Distance integrator | adds up speed into distance | 100 per second |
| Dashboard (laptop) | live display, controls, run table | 50 per second |

# Part 2: The story of the night

Testing ran from about 19:30 on Oct 3 to 02:30 on Oct 4: 58 recorded runs, 36 LiDAR height measurements, 21 dashboard sessions, 3 car drives. Each problem below follows the same shape: what we saw, why it happened, what we changed, and how we know it worked. "Confirmed" means both the code and the data show it; "supported" means the data is consistent with the explanation but it wasn't isolated in a controlled test.

## Problem 1: the app was too slow (supported)

- **Saw:** the early test build processed only 60–120 frames per second instead of 240, and quality dropped as the cart sped up.
- **Why:** it was a Debug build (unoptimised). Fewer frames per second means the ground moves further between frames, so consecutive images overlap less and the match gets worse.
- **Fix:** Release build and 240 fps camera mode. From 20:36 on, every run processed 232–240 frames per second, with each frame taking under a millisecond.
- **Evidence:** within the same setup, quality falls from about 400 to about 60 as the shift per frame grows from 1.4 to 14 pixels, and at a given speed 240 fps means 2.7× less shift than 90 fps. The early builds also differed in height, floor and exposure, so there is no perfectly controlled before/after.

## Problem 2: tilt was counted twice (confirmed)

- **Saw:** speeds that made no sense when the phone was tilted.
- **Why:** LiDAR measures distance along the camera's line of sight. A tilted phone sees the ground further away than its true vertical height. The code stored that slanted distance as the height, and the speed formula then corrected for tilt again. Measured example: 0.132 m slant distance at 38° tilt, real height 0.104 m. With the double correction, speed reads **{p(st['lidar_tilt_example']['speed_error_pct_if_bug'],0)} % too high**.
- **Fix:** convert the LiDAR distance to vertical height exactly once; reject unreliable readings; log every LiDAR frame.
- **Note:** the early runs used a default 0.30 m height (LiDAR wasn't working yet), so no logged run actually used the buggy height. The size of the error comes from geometry plus the logged LiDAR frames.

## Problem 3: the camera crashed and the phone "drove" while sitting still (confirmed)

- **Saw:** after pressing Measure h, camera tracking never came back, and the distance kept climbing with the phone untouched: **{p(st['imu_only_2026-10-03_21-37-55']['distance_m'])} m in 3 s**, and {p(st['imu_only_2026-10-03_21-34-31']['distance_m'])} m in 31 s.
- **Why:** iOS crash logs (14 of them, 20:26–21:43) show the iPhone's camera service crashing. LiDAR leaves the camera in a 30 fps mode, and jumping straight to 240 fps crashed it. With no camera, the filter only had the accelerometer, and accelerometer-only speed drifts without limit.
- **Fix:** restart the camera at 30 fps and step up to 240; a watchdog restarts it if no frames arrive for 1.5 s; and if the camera has been gone for 2 s, any still moment forces speed to zero.
- **Evidence it worked:** the next two height measurements during runs show a 2.1 s camera pause and then recovery, with the phone at rest reading 0.07 m and 0.00 m. No crash reports after 21:43.

## Problem 4: every run read 7 % long (confirmed fix; cause supported)

- **Saw:** three pushes along a 16 ft (4.877 m) taped course read 5.235, 5.246 and 5.214 m: **+{p(min(pre))} to +{p(max(pre))} %**, very repeatable.
- **Ruled out:** wobble, speed and the filter. A slower push had the same error, and the raw camera data added up to the same number.
- **Why:** at 240 fps the iPhone doesn't report its exact lens parameters, so the focal length was estimated from the lens's field-of-view figure. Wide lenses magnify the centre of the image (barrel distortion), exactly where the measurement patch sits.
- **Fix, step 1:** multiply the focal length by 1.073. Three pushes at a new height (23.5 cm) then read {p(min(first))} to {p(max(first))} %, so the error now depended on height.
- **Fix, step 2:** a two-number calibration: focal length ×1.027, plus a 7.9 mm offset because the LiDAR sensor sits higher than the camera lens. All six pushes then land within ±{p(ins['max_abs_pct'],2)} % ({ins['max_abs_m']*1000:.0f} mm on 4.9 m).
- **Independent check:** re-fitting from the raw data gives ×1.0268 and −8.0 mm, essentially identical to the team's values.

## Problem 5: wobble added distance (confirmed behaviour)

- **Saw:** jostling the phone on its mount added 10–15 % distance.
- **Why:** when the phone rocks, the lens really does swing over the floor, and summing speed counts every swing.
- **Fix:** smooth speed over 0.5 s before adding it up (wobbles cancel, steady motion doesn't), plus de-rotation using the gyroscope. On real jostle logs the smoothing removed 21–47 % of the counted distance (the true distance of those tests is unknown, so this shows the effect, not the accuracy).
- **Side effect, also fixed:** the smoothing's correction term switched off abruptly when stopping, making the phone's distance dip 2.5 cm. The dashboard thought that meant a reset and added the whole run again (+4 m). Fixed on both phone and dashboard.

## Problem 6: the first car drive (confirmed)

- **Saw:** at 10 mph the sensor said about 1.5 m/s; at 20 mph it was stuck around 6 m/s.
- **Why 1:** a feature that tried to track changing ride height from how the image expands collapsed from 0.17 m to 0.06 m, so speed read about 0.3× GPS.
- **Why 2:** above about 6 m/s the ground moves more than the correlation window can see in one frame (±64 pixels after downsampling), so the measurement "wraps around" and fails.
- **Fix:** **motion prediction**: shift the search window by the last known motion so it only has to find the change between frames. On the bench this raised the reliable speed from 3 m/s to 9 m/s.

## Problem 7: the second car drive (confirmed)

- **Saw:** good for 40 s (within ±10 % of GPS up to 8 m/s), then 20 mph read like 10–11 mph, and at 30 mph it read zero.
- **Why 1:** the height tracker drifted again (to 0.11 m), so everything read about 0.69× GPS.
- **Why 2 ("zero lock"):** one weak frame at speed made the tracker try "zero motion" as a guess. The camera's own fixed sensor pattern always matches at zero, so it locked onto zero and stayed there.
- **Fix:** height tracking switched off (a car's ride height barely changes, so one LiDAR reading is enough); coast on the last motion for 24 frames before guessing; reject weak results that jump far from the last good motion.

## Problem 8: the final car drive (partly unresolved)

- **Worked:** height steady at 0.171 m; steady cruise at 10–12 mph read **{p(seg[0]['ratio'],3)} × GPS**.
- **Still failed:** around 25–30 mph at night the image blurs (a 1/1000 s exposure at 12.5 m/s smears the image about 65 pixels), quality drops to the noise floor, and the tracker locked onto zero again, reading 0 from t = {vr['zero_lock']['first_t']:.0f} to {vr['zero_lock']['last_t']:.0f} s while GPS said 4–6.6 m/s.
- **Last fix, NOT road-tested:** a "fake stop" rule (an exact zero right after fast motion needs very high confidence) and motion prediction only above 32 px/frame. Both are in the final code, but no car drive was recorded after them.

## What failed and was kept as a finding

- **High mount (79 cm):** read −6.6, −16.7 and −15.0 %. A fix that worked in simulation (subtracting the static image pattern) made it −31 % on the real cart and was reverted. Untested theory: at that height the floor moves only about 1 pixel per frame, and anything static in the image pulls the answer toward zero.
- **Changing ride height mid-run:** −6 % to −25 % with a fixed height; image-expansion tracking gave −3.4 % once and −34 % the next time. It is switched off.
- **Lesson:** the simulator was right about window wrap and blur, and wrong twice about real optics. Only the fully logged real runs exposed that.

# Part 3: How accurate is it, really?

## What we can compare against

There was no lab-grade reference (no wheel encoder, no RTK GPS). There were two references:
- **Distance:** a taped 16 ft (4.877 m) course for the cart. Exact, but only distance, not speed over time.
- **Speed:** the phone's own GPS in the car. Independent of the camera, but noisy, about 1 update per second, and lagging about 0.65 s. iOS reported its own GPS speed uncertainty as about ±2 m/s (median while moving), which is very conservative; the actual jitter measured from the data is about ±0.16 m/s.

## Distance accuracy (cart, taped course)

| Measure (6 runs, 4.877 m course) | Value |
| Before any fix | +{p(min(pre))} to +{p(max(pre))} % |
| Final calibration, same 6 runs | worst {p(ins['max_abs_pct'],2)} % ({ins['max_abs_m']*1000:.0f} mm), typical {p(ins['rmse_pct'],2)} % |
| Final calibration, each run predicted from the other 5 ("leave-one-out") | worst {p(loo['max_abs_pct'],2)} %, typical {p(loo['rmse_pct'],2)} % |
| Range expected for a new run (95 %, statistics on 6 runs) | {p(loo['pi95_new_run_pct'][0])} to +{p(loo['pi95_new_run_pct'][1])} % |
| Monte Carlo, new run, 95 % of the time within | ±{p(S2['abs_p95_pct'],2)} % (about {S2['abs_p95_m']*100:.1f} cm on 4.9 m) |
| Monte Carlo, new run, 99 % of the time within | ±{p(S2['abs_p99_pct'],2)} % |

Why two versions? The calibration was fitted on these same six runs, so checking it on them is a bit like marking your own homework. The leave-one-out numbers are the honest ones. Both say: **about half a percent on a 5 m course**, valid for a 19–24 cm mount on an indoor floor at walking pace.

## Speed accuracy (car, against GPS)

Final drive, between 4.5 and 25 mph, before the overspeed failure (61 GPS readings):
- average error {p(env['bias'],2)} m/s (no meaningful bias), typical error (RMSE) **{p(env['rmse'],2)} m/s**, 95 % of readings within {p(env['p95'],2)} m/s, worst {p(env['max'],2)} m/s;
- steady driving: typical error {p(vr['steady']['rmse'],2)} m/s; accelerating or braking: {p(vr['transient']['rmse'],2)} m/s;
- best steady cruise (10–12 mph): {p(seg[0]['ratio'],3)} × GPS, typical error {p(seg[0]['rmse'],2)} m/s;
- including the zero-lock failure, the typical error is {p(vr['whole_run_including_failures']['rmse'])} m/s, so the failure dominates.

Part of that 0.6 m/s is GPS's own error and lag, especially during acceleration. We can't fully separate the two without a better reference.

## The Monte Carlo simulation, in plain words

A **Monte Carlo simulation** asks: "if we repeated this measurement thousands of times with realistic random errors, how big would the error usually be?" We ran the real filter code 100,000 times for distance and for speed. Each time we fed it fake but realistic sensor readings: the real timestamps and quality scores from actual runs, with random errors whose sizes were measured from the real data (camera jitter, height-measurement repeatability, calibration uncertainty, GPS noise).

Results:
- **Distance:** 95 % of simulated runs within ±{p(S2['abs_p95_pct'],2)} %, 99 % within ±{p(S2['abs_p99_pct'],2)} %. The biggest single cause is the camera-height measurement: LiDAR repeats to about 0.6 mm, which at 19 cm is already 0.3 %.
- **Speed:** 95 % of instantaneous readings within ±{p(VM['abs_p95'],2)} m/s, 99 % within ±{p(VM['abs_p99'],2)} m/s. Averaged over 10 s: 95 % within ±{p(VM['mean_over_10s']['abs_p95'],2)} m/s. The biggest cause is a slow drift of about ±0.5 m/s lasting a few seconds (likely the car pitching and changing camera height, plus GPS lag); per-frame camera noise is tiny by comparison.
- **Check:** the simulation's spread matches the real runs (distance 0.44 % vs 0.44 %; speed 0.52 vs 0.53 m/s once GPS noise is included). That match is partly built in, because the error sizes were measured from the same data. Treat it as a consistency check, not independent proof.

# Part 4: Winning over the judges

## What the judges score (official MHacks 2026 criteria)

The MHacks 2026 Devpost lists four criteria: **Innovation**, **Technical Complexity**, **Usability**, and **Adherence to Theme** (theme: "Come build something that grows"). The Devpost page does not list a dedicated hardware prize, so **confirm with the organizers which track "hardware path" maps to** and whether it has extra criteria. The Devpost submission deadline is **Oct 4, 12:15 pm EDT**.

How this project maps to each criterion:
- **Innovation:** a sensor class that normally costs thousands of dollars, built from a phone, with on-device sensor fusion. The novel bits are phase-correlation flow at 240 fps with motion prediction, LiDAR height with tilt correction, and a quality-weighted Kalman filter.
- **Technical complexity:** real-time FFT image processing in under 1 ms per frame, a 4-state Kalman filter with outlier rejection, a binary telemetry protocol with checksums, Bonjour auto-discovery, a replay tool that reproduces the phone's filter exactly, a synthetic benchmark, and (per the team write-up) 166 automated checks.
- **Usability:** one-button height measurement, live dashboard, works with a phone people already own, every run saved and exportable.
- **Theme ("grows"):** be careful not to force it. Honest angles: the sensor's capability grew run by run through measured fixes (6 m/s → about 10 m/s; 7 % → 0.5 %), and an open, phone-based tool lets student teams grow their testing without expensive equipment. Pick one angle and say it in one sentence.

## The 2-minute pitch (suggested)

1. **Hook (15 s):** "Racing teams pay thousands for a sensor that measures true speed over the ground, because wheels slip. We built one from an iPhone."
2. **How (30 s):** camera at 240 fps watches the ground; LiDAR gives the height; a Kalman filter on the phone fuses camera, motion sensors and GPS; live to the dashboard.
3. **Proof (30 s):** "On a 4.9 m taped course it's within about half a percent after calibration. In a car it tracked GPS within a few percent at steady 10–12 mph, and up to about 20 mph."
4. **The honest story (30 s):** "It started 7 % off and capped at 6 m/s. Every fix came from logged data: a focal-length calibration, a LiDAR tilt bug, motion prediction for the car. We also show what still fails: night blur above about 25 mph and high mounts."
5. **Close (15 s):** "Every run is logged and replayable; the filter is verified against the phone's own output; our error bounds come from 100,000 simulated runs."

## The live demo (what to show, what not to)

- **Do:** push the cart along a taped line (ideally the 16 ft course at about 19–24 cm mount height), press Measure h first, show the dashboard live, then the runs table with the distance and the error against the tape.
- **Do:** show the quality (PSR) dropping when you cover the lens or lift the phone. It proves the system knows when it can't see.
- **Do:** have the car-drive plot ready (estimated speed vs GPS over time) instead of promising a live car demo.
- **Don't:** mount it at a new height without measuring h; don't demo on a shiny or featureless floor; don't claim 30 mph.
- **Have ready:** the slides (Google Slides in your Drive), a phone screenshot of a clean run, and the GitHub repo link.

## Questions judges are likely to ask, with answers

### About the idea

**"Why not just use GPS?"** Phone GPS updates about once a second, lags about 0.65 s, and its speed is noisy, especially at low speed. It can't measure wheel slip. The camera measures 240 times a second and works indoors.

**"Why not wheel speed sensors?"** They measure the wheel, not the ground. When tyres slip, which is exactly what vehicle-dynamics testing studies, wheel speed is wrong. A ground-speed sensor is how you measure slip.

**"How does this compare to a Correvit?"** Same principle (optical ground speed), far cheaper hardware. Ours is validated only at low speed, short distances and low mounts, and is blur-limited at night. Don't claim parity; claim the core function for a fraction of the cost.

**"Isn't this just an optical mouse?"** Same principle, very different problem: the ground is centimetres to metres away (so height must be measured), it moves up to hundreds of pixels per frame (needs motion prediction), lighting and texture vary, and the phone vibrates and tilts (needs gyro correction and outlier rejection).

### About the camera / optical flow

**"What is phase correlation and why use it?"** It compares two images in the frequency domain (via FFT) and produces a sharp peak at their relative shift. It gives sub-pixel shifts, costs the same for every frame (well under 1 ms on the phone), and comes with a built-in confidence score (PSR). Feature tracking (finding corners) struggles on low-texture surfaces like concrete at night.

**"What's PSR?"** Peak-to-sidelobe ratio: how much the best match stands out from the background. Noise sits around 3–7; good texture gives tens to hundreds. Below 8 we drop the frame; above that, the filter trusts the frame in proportion to PSR squared.

**"What limits top speed?"** Two things. (1) Window wrap: without prediction the search window covers ±64 downsampled pixels per frame, about 6 m/s at a 17 cm mount at 240 fps. Motion prediction removes that. (2) Motion blur: at night with a 1/1000 s exposure, 12.5 m/s smears the image about 65 pixels and tracking fails. Daylight (shorter exposure) or a higher mount extends the range.

**"Why 240 fps?"** At a given speed, the ground moves 2.7× less per frame than at 90 fps, so consecutive images overlap more and the match is better. In our data quality falls steeply as the shift per frame grows.

**"Why does height matter so much?"** The camera measures pixels; pixels become metres by multiplying by height ÷ focal length. A 1 % height error is a 1 % speed and distance error.

### About the Kalman filter (expect the deepest questions here)

**"What are your states?"** Forward speed, sideways speed, and the accelerometer bias on each axis. Distance is computed outside the filter by adding up speed.

**"What is the prediction step / process model?"** Speed changes by the measured acceleration minus the estimated bias, with a term for turning (yaw rate couples forward and sideways speed). The bias is modelled as slowly wandering.

**"What do Q and R mean, and how did you choose them?"** Q is how much we expect the model to be wrong each step (process noise); R is how noisy each measurement is. Q for acceleration was raised from 0.02 to 0.1 after the first real run because the cart's jolts were being rejected as outliers. R for the camera is 1.6e-3 × (20/PSR)², so a sharp frame is trusted much more than a weak one. GPS uses its own reported accuracy squared as R.

**"Why a Kalman filter and not just averaging?"** It weights each measurement by how reliable it is right now, fills gaps (the IMU carries the estimate between camera frames and through brief dropouts), estimates and removes accelerometer bias, and reports its own uncertainty.

**"How do you handle outliers?"** Each camera reading is compared to the prediction. If it is more than 3 standard deviations away (a test statistic called NIS above 9) it is rejected. If 12 in a row are rejected, the filter assumes it is the one that diverged, widens its uncertainty and accepts the camera again.

**"What is ZUPT?"** A zero-velocity update. When the IMU says the phone is perfectly still, the filter is told the speed is exactly zero. This kills drift at every stop.

**"Why only use GPS above 1 m/s?"** GPS gives speed magnitude, and the filter needs a direction to apply it; near zero the direction is undefined and GPS speed noise is relatively huge. Honest follow-up: in our tuning GPS barely affects the output (at most 4 mm/s), because the camera is trusted far more. Its real value would be as a sanity check to catch camera failures like the zero lock. That's a next step.

**"Did you verify the filter?"** Yes. A replay tool feeds the logged sensor data back through the filter and reproduces the phone's own output; our independent Python port matched it on all 53 runs from 20:06 on (typical difference 0.000002 m/s). The repo also has a synthetic test run with a known 10.000 m answer and, per the write-up, 166 automated checks (not re-run in this analysis, which had no Swift toolchain).

### About accuracy and testing

**"How accurate is it?"** Distance: about ±0.5 % on a 5 m taped course (worst leave-one-out run 0.54 %; 95 % bound from simulation ±{p(S2['abs_p95_pct'],1)} %). Speed in the car: typical error {p(env['rmse'],2)} m/s against GPS between 4.5 and 25 mph, about 0.27 m/s in steady cruising. Always say what it was tested on: indoor floor, 19–24 cm mount, walking pace for distance; one night drive for speed.

**"How did you get ground truth?"** A taped course for distance; phone GPS for car speed. We don't have a wheel encoder or RTK GPS, so speed accuracy partly includes GPS's own error. That's the first thing we'd add.

**"Isn't calibrating on six runs and testing on the same six cheating?"** Yes, partly, which is why we also report leave-one-out (each run predicted from the other five): worst error 0.54 % instead of 0.32 %. Both calibration numbers have physical meaning (lens magnification and the LiDAR-to-lens offset), not just curve fitting.

**"What's your error budget?"** Distance is dominated by camera-height repeatability (LiDAR repeats to about 0.6 mm, which is 0.3 % at 19 cm). Per-frame camera noise contributes under 0.1 %. Speed in the car is dominated by a slow drift of about ±0.5 m/s, likely ride height changing as the car pitches, plus GPS lag in the reference.

**"What's a Monte Carlo and why did you do it?"** We ran the actual filter 100,000 times with realistic random errors (sizes measured from our data) to get error bounds that include all the error sources together, including the ones that don't add up simply. It agrees with the real runs.

### About the process and decisions

**"What was the hardest bug?"** Good answers: the zero lock (the fixed sensor pattern looks like a perfect "no motion" match), or the 7 % focal error (we ruled out wobble, speed and the filter before finding the lens). Tell it as symptom → evidence → cause → fix.

**"Why turn off the height tracker?"** It worked on synthetic floors and roughly on the cart, but drifted badly in both car drives (0.17 → 0.06 m, then 0.11 m). A car's ride height changes little, so one LiDAR reading at the start beats a drifting estimate.

**"Why does the high mount fail?"** Honestly: not proven. Leading theory: at 79 cm the floor moves only about 1 pixel per frame at walking pace, and static content in the image pulls the measured shift toward zero. Our attempted fix made it worse on real hardware, so we reverted and limited the claimed operating range.

**"Is it safe to test in a car?"** Be ready to say who drove, where it was mounted, and that the phone was rigidly attached. Don't improvise.

**"Who did what?"** From the README: app and phone pipeline (David), Kalman filter (Joseph), dashboard (Sean). Each person should be able to explain their part in two sentences and the whole system in one.

## Claims to avoid (where the data disagrees with the write-up)

These are places where the team's write-up says something the logged data doesn't fully support. A judge who reads the repo could catch them.

| Don't say | Say instead |
| "All six runs within ±0.4 %" (as a general accuracy) | "Within ±0.4 % on the calibration runs; about ±0.5 % when each run is predicted from the others." |
| "Runs 15–17 at 79 cm read 15–17 % low" | "Two read 15–17 % low; one faster push read −6.6 %." |
| "Phone GPS speed is good to ±0.3–0.5 m/s" | "Measured GPS jitter is about ±0.16 m/s, but iOS reports ±2 m/s and it lags about 0.65 s." |
| "Consistent readings up to 25 mph" | "Tracks GPS well to about 20 mph; near 25 mph it reads 10–20 % low at night, and above that it lost tracking." |
| "The zero-lock is fixed" | "We wrote a fix, but haven't driven it yet." |
| "GPS fusion improves accuracy" | "GPS is fused, but the camera dominates; GPS's real role should be a failure check." |
| Any accuracy number without its conditions | "…on an indoor floor, 19–24 cm mount, walking pace" / "…in one night drive, against phone GPS" |

## Rules for answering when you don't know

- Say what the data shows, then what you think, and label which is which: "We measured X. We think it's because Y, but we didn't isolate it."
- If you don't know: "We didn't test that. Here's how we would." That is far better than guessing.
- Show failures before you're asked. Judges reward understanding, not perfection.

# Part 5: Cheat sheet

## Numbers to know

| Thing | Number |
| Camera rate / processing time | 240 fps / under 1 ms per frame |
| Correlation window | 128 × 128 (from a 256 × 256 crop, downsampled 2×) |
| Distance accuracy (cart, 4.9 m) | ±0.5 % (worst cross-validated 0.54 %; MC 95 % ±{p(S2['abs_p95_pct'],2)} %) |
| Before calibration | +7 % |
| Speed vs GPS (car, 4.5–25 mph) | RMSE {p(env['rmse'],2)} m/s; steady cruise {p(seg[0]['ratio'],3)} × GPS |
| Speed MC bound (instantaneous) | ±{p(VM['abs_p95'],2)} m/s (95 %), ±{p(VM['abs_p99'],2)} m/s (99 %) |
| Top speed without prediction | about 6 m/s at a 17 cm mount |
| Night blur limit | about 25 mph (11 m/s) at 1/1000 s exposure |
| Calibration | focal ×1.027, LiDAR −7.9 mm |
| Filter tuning | q_accel 0.1, q_bias 1e-5, R_flow 1.6e-3·(20/PSR)², gate 9 (3σ), PSR min 8 |
| GPS lag / jitter | 0.65 s / about 0.16 m/s |
| Data recorded | 58 runs, 36 LiDAR logs, 3 car drives |

## Glossary

- **Optical flow:** how much an image pattern moves between frames.
- **Phase correlation:** an FFT-based way to measure that shift in one step.
- **PSR (peak-to-sidelobe ratio):** confidence of the match; noise is about 3–7.
- **Focal length (in pixels):** how zoomed-in the lens is; with height, it converts pixels to metres.
- **LiDAR:** a laser rangefinder on the iPhone's back; gives camera height.
- **IMU:** inertial measurement unit: accelerometer (acceleration) plus gyroscope (rotation rate).
- **Kalman filter:** blends a model's prediction with measurements, weighted by uncertainty.
- **State:** what the filter estimates (here: two speeds, two accelerometer biases).
- **Q / R:** how uncertain the model is / how noisy a measurement is.
- **Gating / NIS:** rejecting measurements too far from the prediction (more than 3σ).
- **ZUPT:** zero-velocity update; "you're stopped, speed is zero".
- **Bias:** a constant offset error in a sensor.
- **Motion prediction:** moving the search window by the expected shift so fast motion stays in view.
- **Window wrap:** a shift too big for the window wraps around and reads wrong.
- **Barrel distortion:** wide lenses magnify the image centre relative to the edges.
- **Leave-one-out:** test each run using a calibration fitted without it.
- **RMSE:** root-mean-square error, the typical size of an error.
- **95 % bound:** 95 % of measurements fall inside it (not the same as a guarantee).
- **Monte Carlo:** repeating a simulated measurement many times with random errors to see the spread.
"""


pv = mcv["params"]
dvd = mcd["variance_decomposition"]
TECH = f"""
# Part 6: Technical deep dive

> This part is for learning the system properly: how the sensors are fused, what each sensor is for and where its hardware limits are, how the Kalman filter works and why each parameter has the value it has, and how the Monte Carlo simulation was built. Every constant below is quoted from the code (GroundSpeedKit and the iOS app) or measured from the logs.

## 6.1 Sensor fusion: how it actually works

### The big idea

No single sensor is good enough on its own:
- the **camera** measures speed directly but only when it can see texture, and only as well as the height and focal length are known;
- the **IMU** reacts instantly and never "loses sight", but drifts within seconds;
- **GPS** never drifts but is slow, laggy and noisy;
- **LiDAR** gives the height that turns camera pixels into metres, but can't run while the camera is filming at 240 fps.

Fusion means letting each sensor do what it is good at and covering its weakness with another. In this system the division of labour is:

| Job | Done by | Why that sensor |
| Measure speed over ground | camera (optical flow) | direct measurement of motion, 240 times a second |
| Convert pixels to metres | LiDAR height h (once) + focal length f | sets the scale of every measurement |
| Fill in between camera frames and through short dropouts | IMU accelerometer (filter predict step) | 100 Hz, never blind |
| Detect "stopped" and pin speed to zero | IMU (stillness) + camera (near-zero flow) | kills drift at every stop |
| Remove fake motion from the phone rotating | IMU gyroscope (de-rotation) | rotation sweeps the image without the vehicle moving |
| Correct for a tilted phone | IMU gravity vector | gives the tilt angle |
| Slow, independent speed check | GPS | absolute, never drifts (but weak in this tuning) |

### What happens to one sample, step by step

1. **Everything is timestamped on one clock** (the phone's monotonic clock), then put through a 50 ms reorder buffer so samples are processed strictly in time order. At equal timestamps the order is always IMU, then camera, then GPS, then zero-velocity. Every input is rounded to exactly what gets written to the log, so replaying the log reproduces the phone's output bit-for-bit. That's why we could verify the filter offline.
2. **IMU sample (100 Hz):** gravity is already removed by iOS; the calibrated bias (measured while still) is subtracted. The filter runs its **predict** step. The distance integrator adds the new speed. The stillness detector decides whether a zero-velocity update is due.
3. **Camera sample (240 Hz):** the phase-correlation shift becomes a velocity using h and f. Then **de-rotation**: the average gyro rate since the previous frame, times h, is subtracted (forward speed minus h·ω_y, sideways speed plus h·ω_x). If the phone was rotating, the frame's quality is lowered (only about half of the rotation effect is removed in practice, so the leftover 0.5·h·|ω| is treated as extra noise). Then the filter runs its **update** step, weighted by PSR.
4. **GPS fix (~1 Hz):** used as an update only if both GPS speed and the filter's speed are above 1 m/s.
5. **Zero-velocity update:** applied after everything else at that timestamp, if the vehicle is judged still.

### Who really controls the output?

With the chosen numbers the filter follows the camera very closely. At typical quality it settles onto a new camera reading in about 5–20 ms:

| Camera quality (PSR) | Camera noise assumed (1σ) | Filter's own uncertainty after update | Filter time constant |
| 10 (weak) | 0.080 m/s | 0.038 m/s | ~19 ms |
| 20 | 0.040 m/s | 0.025 m/s | ~11 ms |
| 50 (typical car) | 0.016 m/s | 0.013 m/s | ~6 ms |
| 100+ (good floor) | 0.008 m/s | 0.008 m/s | ~5 ms |

So in practice: **the camera sets the speed, the IMU bridges the gaps** (there are 240 camera frames but only 100 IMU steps per second, plus dropped frames and pauses), and **GPS barely matters** (removing it changes the output by at most 4 mm/s). That also explains the biggest weakness: if the camera is confidently wrong (the zero lock, PSR about 13, above the cut-off of 8), nothing in the filter can overrule it.

### Axes and sign conventions

The vehicle frame is x forward, y left, z up. Camera axes (x right, y down in the image) are mapped to vehicle axes by a mount setting. One finding from the data: every calibration push actually moved along the phone's **−x** axis (the summed forward speed was −4.9 m). It doesn't matter for the odometer, which adds up the magnitude of speed, but the signed "net forward" field read negative on those runs.

## 6.2 Each sensor: purpose and hardware limits

### Camera (the primary sensor)

- **Configuration:** wide camera, 1280×720 at 240 fps (120 fps if the phone gets hot), video stabilisation off, focus locked after a 1.5 s settle with the torch on, exposure capped at 1/1000 s with ISO raised to compensate (it ran at ISO 2200 at night). A 256×256 centre crop is downsampled to 128×128 for the correlator.
- **Purpose:** measures the ground's shift between frames plus a quality score (PSR).
- **Limit: texture and light.** A blank, shiny or dark surface gives low PSR. Frames with PSR below 8 are skipped.
- **Limit: motion blur.** Blur in pixels ≈ speed × focal length × exposure ÷ height. At 12.5 m/s, 17 cm height and 1/1000 s that is about 65 px, and tracking fails. Shorter exposures need more light: halving it at night gave no net gain.
- **Limit: window wrap.** The 128-pixel window (after 2× downsampling) can only see shifts up to ±64 px per frame, about 6 m/s at a 17 cm mount. Motion prediction (offsetting the window by the last motion; used above 32 px/frame in the final build) removes this limit.
- **Limit: fixed pattern.** The sensor's own static noise always correlates perfectly at zero shift (PSR 10–25). When the real match is weak, the tracker can lock onto zero. Mitigations: coast 24 frames on the last motion, require PSR ≥ 30 to accept a big jump, and (final build, untested) require PSR ≥ 30 for a sudden exact zero after fast motion.
- **Limit: no lens data at 240 fps.** iOS doesn't deliver per-frame intrinsics in this mode, so the focal length is computed from the field-of-view figure (73.3°) and calibrated (×1.027).
- **Limit: height range.** Below about 20 cm the lens struggles to focus and the window covers little ground; at 79 cm the floor moves only about 1 px per frame at walking pace and readings were 7–17 % low.
- **Limit: the iOS camera service.** Switching straight from the LiDAR's 30 fps mode to 240 fps crashed it (14 crash reports). Fix: start at 30 fps and step up; a watchdog restarts the camera after 1.5 s without frames.
- **Limit: heat and dropped frames.** Under thermal pressure the app drops to 120 fps, which halves the speed range. Logged frame drops were about 1 % in the final car drive and about 4 % in the first.

### LiDAR (sets the scale)

- **Purpose:** measures camera height h before a run ("Measure h"), which converts pixels to metres. It also zeroes the IMU bias at the same still moment.
- **How:** a separate capture session pauses the flow camera for about 1.5 s, reads 320×180 depth maps, takes the median of the centre patch of each frame, then the median across frames.
- **Limit: can't run with the 240 fps camera.** iOS can't run both sessions at once, so height is measured once, not tracked. Changes in ride height during a run (suspension, pitch) are invisible. That is the leading suspect for the ±0.5 m/s slow error in the car.
- **Limit: it measures along its line of sight.** On a tilted phone that is longer than the vertical height, so the reading is multiplied by cos θ from the gravity vector. Getting this wrong was the +61 % tilt bug.
- **Limit: it isn't at the lens.** The LiDAR sits about 7.9 mm higher than the camera's optical centre, hence the −7.9 mm calibration offset (fit uncertainty about ±2.3 mm).
- **Accepted range and checks:** 5 cm to 2 m; reading rejected if the depth isn't metric, under half the pixels are valid, or the frames disagree by more than 8 % or 1 cm. It reads unstably below about 8 cm.
- **Repeatability:** repeated measurements at an unchanged mount scatter by about 0.6 mm (1σ), and the app stores h rounded to 1 mm. At a 19 cm mount that is already about 0.3 % in distance, which is most of the run-to-run scatter we see.

### IMU: accelerometer and gyroscope (bridging, stillness, rotation, tilt)

- **Configuration:** iOS Core Motion "device motion" at 100 Hz. iOS's own fusion removes gravity and provides the gravity direction. The app rotates samples into the vehicle frame and subtracts a bias measured during a 2 s still calibration (rejected if the phone moved: accel standard deviation above 0.04 m/s²).
- **Purpose 1, prediction:** carries the speed estimate between camera frames and through dropouts.
- **Purpose 2, stillness:** declared still when the variation of acceleration magnitude over 0.3 s is below 0.0025 (m/s²)² (a standard deviation of 0.05 m/s²) and the latest camera speed is below 0.03 m/s. A zero-velocity update follows.
- **Purpose 3, de-rotation:** gyroscope rate × h is subtracted from the camera speed.
- **Purpose 4, tilt:** the gravity direction gives the camera's tilt for the cos θ correction.
- **Limit: drift.** Integrated acceleration drifts fast: with no camera the phone "drove" 15 m in 3 s while sitting still. That's why, once the camera has been gone for 2 s, any still moment forces speed to zero.
- **Limit: noise and vibration.** At rest the accelerometer noise was about 0.046 m/s² on the cart and 0.062 m/s² in the car; the car mount saw spikes of about 3–4 g. Large jolts are why q_accel was raised.

### GPS (independent but weak)

- **Configuration:** Core Location "best for navigation", every fix delivered; only the speed and its reported accuracy are used (position isn't).
- **Purpose:** an absolute speed that never drifts; used above 1 m/s.
- **Limit: slow and late.** About 1 fix per second, trailing the camera by about 0.65 s (measured by cross-correlation in two drives).
- **Limit: noise.** Measured jitter about 0.16 m/s, but iOS reports a 1σ "speed accuracy" of about 2.2 m/s (median while moving in the final drive). The filter uses the reported value squared as R, so GPS gets very little weight.
- **Limit: no use indoors**, and near zero speed its direction is meaningless.
- **Note:** GPS position accuracy (±5 m) was a red herring early on. Only speed is used.

### Torch, Wi-Fi and the dashboard

- **Torch:** lights the ground so short exposures are possible; at night the camera still needed ISO 2200.
- **Wi-Fi telemetry:** 50 updates a second over UDP to the laptop (48- or 72-byte frames with a checksum), commands over TCP, Bonjour discovery. In the car up to about 2,300 packets were lost in one session. That doesn't affect accuracy because the phone records everything locally; the dashboard is display only.

## 6.3 The Kalman filter in detail

### State and model

- **State (4 numbers):** forward speed v_x, sideways speed v_y (m/s), and accelerometer biases b_x, b_y (m/s²). Position and distance are **not** states; distance is integrated outside the filter so every filter variant shares the same odometer.
- **Predict (every IMU sample, Δt ≈ 10 ms, clamped to 0.1 s):** v_x ← v_x + Δt·(a_x − b_x + r·v_y) and v_y ← v_y + Δt·(a_y − b_y − r·v_x), where r is the yaw rate from the gyroscope (turning swaps forward and sideways speed). Biases are assumed constant apart from a slow random walk.
- **Uncertainty grows each predict** by Q = diag(q_accel·Δt, q_accel·Δt, q_bias·Δt, q_bias·Δt).
- **Updates are scalar and sequential** (forward axis, then sideways). Each needs one division instead of a matrix inverse, and the covariance uses the **Joseph form** (more robust to rounding) and is forced symmetric.

### Parameters and why they have these values

| Parameter | Value | Meaning | Why this value |
| q_accel | 0.1 (m/s²)²·s | how much speed can change unpredictably per second (1σ ≈ 0.32 m/s in 1 s) | raised from 0.02 after the first on-device run: cart jolts were being rejected as outliers; 0.1 halved the rejected camera frames |
| q_bias | 1e-5 | how fast the accelerometer bias can wander (≈ 0.03 m/s² in 100 s) | bias is nearly constant over a run; small keeps the bias estimate stable |
| r_flow_base | 1.6e-3 (m/s)² | camera noise variance at the reference quality | camera σ = 0.04 m/s at PSR 20 |
| psr_ref | 20 | reference PSR for that noise | R scales as (20/PSR)²: double the PSR, a quarter of the variance |
| psr_min | 8 | below this, frames are skipped, not down-weighted | the correlator's noise floor is about PSR 3–7; a covered lens reports a confident zero, so weak frames must be ignored, not merely trusted less |
| gate | 9 | reject a camera sample if its normalised innovation squared (NIS) exceeds 9 on either axis | 9 = 3σ |
| gate_reset_count | 12 | after 12 rejected camera samples in a row, reopen the filter's uncertainty and accept | if the filter itself diverged (e.g. after IMU-only drift), endless rejection would lock it out forever |
| r_zupt | 1e-4 (m/s)² | zero-velocity "measurement" noise (σ 0.01 m/s) | strong pull to zero when still; never gated |
| p0_v / p0_b | 0.25 / 0.01 | starting uncertainty: speed σ 0.5 m/s, bias σ 0.1 m/s² | wide enough to lock on quickly |
| GPS R | speed_acc² | GPS noise from iOS's own estimate | honest, but iOS's figure is very conservative, so GPS gets little weight |
| GPS threshold | > 1 m/s (both GPS and filter) | GPS speed is a magnitude; the update needs a direction | near zero the direction is undefined and GPS noise is relatively huge |

### Rules that aren't obvious

- **The gate uses the prediction from before the update,** on both axes. If either axis fails, the whole camera sample is rejected.
- **The gate is switched off after start-up** until the first accepted measurement, so a filter started while already moving can lock on instead of rejecting everything.
- **Lockout recovery:** on the 12th straight rejection, the cross-terms of the speed uncertainty are cleared, the speed variances are raised to at least 0.25, and the sample is accepted.
- **The GPS update is linearised** (an "extended" Kalman update): the measurement is |v|, so the update direction is the current speed direction v/|v|.
- **Zero-velocity updates are never gated** and always applied after everything else at that timestamp.

### Design decisions and their trade-offs

- **Why 4 states and not more?** Speed plus bias is the minimum that removes accelerometer drift. Scale (h, f) is not a state: without a reference that sees true speed, scale isn't observable from these sensors, and GPS is too weak to estimate it. Consequence: scale errors pass straight through. The filter's own σ (about 0.01–0.04 m/s) is far smaller than the real error (about 0.5 m/s in the car), so **the filter is overconfident about slow, systematic errors.**
- **Why skip low-quality frames instead of down-weighting?** A covered or blank lens returns zero shift with a modest PSR; down-weighting would still pull speed toward zero.
- **Why a 4-state coupled filter rather than two separate axes?** Turning couples the axes through the yaw rate. A simpler alternative, DecoupledKF2x2 (two independent 2-state filters treating the coupling as a known input), is also in the code and selectable, but ReferenceKF4 is the default.
- **Why distance outside the filter?** All filter variants share one odometer, and it can include the wobble smoothing (0.5 s low-pass), the 0.05 m/s deadband and the lead compensation without touching the filter.
- **Why deterministic replay?** Quantising inputs to the logged values and fixing the processing order means `kfreplay` (Swift) and our Python port reproduce the phone's output exactly. Every tuning change can be tested on real recorded runs before going back on the phone.
- **What we'd change next:** add a consistency check between camera and GPS (or IMU) to detect a confidently wrong camera such as the zero lock; model camera noise as speed- and vibration-dependent; consider a scale or height state if a better reference (wheel encoder, RTK) is added.

### The odometer (distance)

- Integrates the magnitude of the smoothed speed by the trapezoid rule after each IMU step (Δt clamped to 0.1 s).
- **Low-pass τ = 0.5 s:** wobble back-and-forth averages out; steady motion is only delayed.
- **Lead term τ·|v̄|:** added to the displayed distance to remove that 0.5 s delay while moving; it fades to zero when stopping (an abrupt cut-off caused the 2.5 cm dip and the dashboard's +4 m bug).
- **Deadband 0.05 m/s:** slower speeds count as zero. On the calibration pushes the deadband and smoothing together lose about 0.5 % of distance, and the focal calibration absorbs that. It depends slightly on how a push accelerates (about ±0.17 % between pushes).

## 6.4 The Monte Carlo simulation in detail

### Why simulate at all?

The real tests give six distance runs and one usable car drive: too few to state a 99 % bound directly. The pipeline is also non-linear (skipped frames, gating, deadband, magnitude of a 2-D speed, quality-weighted noise) and its errors are correlated in time, so simple "add the variances" formulas aren't reliable. A Monte Carlo runs the actual pipeline many times with realistic random errors and reads the bounds off the results.

### Design decisions

- **Run the real algorithm.** The filter and odometer were ported to Python and vectorised so 100,000 trials run side by side. The port reproduces the phone's own output (typical difference 2e-6 m/s), and the vectorised version matches the scalar port to 1e-16 m/s.
- **Use real conditions.** Every trial uses the real timestamps, camera quality sequence and stillness events of a recorded run. Only the sensor errors are random.
- **Define the truth.** Distance: each of the six calibration pushes' own speed profile, smoothed over 0.5 s, scaled so its forward displacement is exactly 4.8768 m (a straight course). Speed: the first 95 s of the final car drive (steady cruise, acceleration to 8.7 m/s, braking), smoothed.
- **Measure the error sizes from data, don't invent them.** Each input below comes from the logs; the two exceptions are labelled as assumptions.
- **Keep the four quantities separate:** target (true), measured (camera), estimated (filter output), and the Monte Carlo distribution of estimated minus target.

### Inputs and where each came from

| Input | Value used | Source |
| Camera frame noise | drawn from the real residuals in the same PSR band (robust σ ≈ {M['noise']['flow_cart_robust_sd']*1000:.0f} mm/s cart, {M['noise']['flow_car_robust_sd']*1000:.0f} mm/s car); car noise correlated frame to frame (ρ = {pv['rho_white_240hz']:.2f}) | residual of each frame from a 0.1 s running median; tails are heavy (excess kurtosis 15–25), so we resample the real residuals instead of assuming a bell curve |
| Camera height h | σ = {dvd['height_sd_m_incl_rounding']*1000:.2f} mm | repeated LiDAR measurements at an unchanged mount (0.60 mm) plus 1 mm rounding in the app |
| Calibration (focal ×c, offset δ) | c = {M['calibration_refit']['focal_correction']:.4f} ± {M['calibration_refit']['focal_correction_se']:.4f}, δ = {M['calibration_refit']['height_offset_m']*1000:.1f} ± {M['calibration_refit']['height_offset_se_m']*1000:.1f} mm, correlation {M['calibration_refit']['param_corr']:.2f} | least-squares refit on the six runs; drawn as a correlated pair because the two trade off against each other |
| Unexplained per-run error (distance) | {dvd['unexplained_sd_pct']:.2f} % | variance matching: height and camera noise already explain all of the observed {dvd['empirical_residual_sd_pct']:.2f} % scatter |
| Slow speed error (car) | σ = {pv['sd_slow_abs']:.2f} m/s, correlation time {pv['tau_slow_s']:.1f} s (a random process that wanders and returns to zero) | car-vs-GPS errors at PSR ≥ 15 minus GPS's own jitter; correlation from consecutive 1 s errors |
| GPS (car) | jitter {pv['sd_gnss']:.2f} m/s, lag {abs(pv['lag_s']):.2f} s | second difference of the 1 Hz GPS series; cross-correlation with the camera |
| IMU | noise {pv['sd_imu']:.3f} m/s² (car), 0.046 (cart); bias σ 0.02 m/s² | noise from still periods; **bias σ is an assumption** (it barely matters because the camera dominates) |
| Trials | 100,000 each | fixed random seeds (20261004 and +7) so results are reproducible |

### Decisions made along the way (and why)

- **Smoothed truth.** Using the raw (noisy) filter speed as truth inflated the "true" distance, because the magnitude of a noisy 2-D speed is biased upward. Smoothing fixed that.
- **Normalising the pipeline bias.** The pipeline alone reads about {p(mcd['pipeline_pilot']['raw_pipeline_bias_pct'],2)} % on these pushes (deadband and smoothing). The real calibration was fitted on the real pipeline, so it already absorbs that; the simulation divides it out once, like the calibration did. The push-to-push spread of that bias stays in as a real effect.
- **Additive, not proportional, slow error.** A slow error proportional to speed predicted a spread of 0.79 m/s against the real 0.53 m/s, because real segment errors didn't grow with speed. The additive model matches (0.52 m/s), so it is the main model; the proportional one is kept as a sensitivity case.
- **Two distance scenarios.** "Repeat run" keeps the calibration fixed (comparable with the in-sample errors); "new run" also randomises the calibration constants (comparable with leave-one-out). The new-run numbers are the ones to quote.
- **Random subsets for convergence.** Trials are generated push by push, so the first N trials all come from the first pushes. An early convergence check used those prefixes and looked lopsided; it was redone with random subsets.

### Results and checks

- **Distance (new run):** 95 % within {p(S2['int95_pct'][0],2)} to +{p(S2['int95_pct'][1],2)} %; 99 % within {p(S2['int99_pct'][0],2)} to +{p(S2['int99_pct'][1],2)} %; |error| below {p(S2['abs_p95_pct'],2)} % ({S2['abs_p95_m']*100:.1f} cm) 95 % of the time.
- **Speed (car, instantaneous):** 95 % within ±{p(VM['abs_p95'],2)} m/s, 99 % within ±{p(VM['abs_p99'],2)} m/s; 10 s averages within ±{p(VM['mean_over_10s']['abs_p95'],2)} m/s 95 % of the time.
- **Convergence:** the 95th and 99th percentiles change by less than a few hundredths of a percent from 10,000 to 100,000 trials, and five extra random seeds agree.
- **Analytic check:** adding the independent error terms by hand gives {p(mcd['analytic_sd_pct'],2)} % against the simulation's {p(S2['sd_pct'],2)} %.
- **Against reality:** distance spread {p(mcd['empirical_vs_mc']['loo_vs_S2']['emp_sd'],2)} % real vs {p(mcd['empirical_vs_mc']['loo_vs_S2']['mc_sd'],2)} % simulated; car speed {p(mcv['empirical_vs_mc']['real_sd'],2)} vs {p(mcv['empirical_vs_mc']['mc_obs_sd'],2)} m/s (after adding GPS jitter), with matching 95th percentiles.
- **Sensitivity:** swapping the camera-noise model (real residuals, bell curve, heavy-tailed) changes almost nothing. Tripling camera noise widens distance bounds only slightly. The bounds are dominated by height and calibration (distance) and by the slow drift (speed). Extrapolating the calibration to a 0.50 m mount widens the distance 95 % interval to about ±1.5 %.

### What the Monte Carlo cannot tell you

- **It's partly circular.** The slow speed error and the per-run scatter were measured from the same runs the simulation is compared with, so matching spreads are expected; the matching tails and time behaviour are the real check.
- **It doesn't simulate failures:** blur above about 25 mph, the zero lock, high mounts, or ride-height changes. The bounds only apply inside the tested envelope.
- **It inherits the reference's limits:** six runs on one floor at two heights, and one night drive measured against phone GPS.
- **Its numbers are prediction intervals** (where 95 % of individual measurements fall), not confidence intervals of an average.
"""
TEXT = TEXT + TECH


def to_html(text):
    """Same mini-syntax → HTML (for pasting into an existing Google Doc)."""
    out, lines, i, para = [], text.strip().split("\n"), 0, []
    b = lambda s: re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

    def flush():
        if para:
            out.append(f"<p>{b(' '.join(para))}</p>"); para.clear()
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln:
            flush(); i += 1; continue
        for pre, tag in (("### ", "h3"), ("## ", "h2"), ("# ", "h1")):
            if ln.startswith(pre):
                flush(); out.append(f"<{tag}>{b(ln[len(pre):])}</{tag}>"); break
        else:
            if ln.startswith("- ") or ln[:3] in {f"{n}. " for n in range(1, 10)}:
                flush(); kind = "ul" if ln.startswith("- ") else "ol"; items = []
                while i < len(lines) and (lines[i].startswith("- ") if kind == "ul" else lines[i][:3] in {f"{n}. " for n in range(1, 10)}):
                    items.append(f"<li>{b(lines[i][2:] if kind == 'ul' else lines[i][3:])}</li>"); i += 1
                out.append(f"<{kind}>{''.join(items)}</{kind}>"); continue
            elif ln.startswith("> "):
                flush(); out.append(f'<table border="1" style="border-collapse:collapse;width:100%"><tr><td style="background:#EAF2FB;padding:6px">{b(ln[2:])}</td></tr></table><p></p>')
            elif ln.startswith("| "):
                flush(); rows = []
                while i < len(lines) and lines[i].startswith("| "):
                    rows.append([x.strip() for x in lines[i].strip().strip("|").split(" | ")]); i += 1
                h = "".join(f'<th style="background:#0072B2;color:#ffffff;padding:4px">{b(c)}</th>' for c in rows[0])
                body = "".join("<tr>" + "".join(f'<td style="padding:4px">{b(c)}</td>' for c in r) + "</tr>" for r in rows[1:])
                out.append(f'<table border="1" style="border-collapse:collapse"><tr>{h}</tr>{body}</table><p></p>'); continue
            else:
                para.append(ln.strip())
        i += 1
    flush()
    return "\n".join(out)


def add_runs(par, text, size=None):
    for i, part in enumerate(re.split(r"(\*\*[^*]+\*\*)", text)):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        r = par.add_run(part[2:-2] if bold else part)
        r.bold = bold
        if size:
            r.font.size = Pt(size)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcolor)
    tcPr.append(s)


def main():
    doc = Document()
    stl = doc.styles["Normal"]; stl.font.name = "Calibri"; stl.font.size = Pt(11)
    lines = TEXT.strip().split("\n")
    i = 0
    para_buf = []

    def flush():
        if para_buf:
            add_runs(doc.add_paragraph(), " ".join(para_buf))
            para_buf.clear()

    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln:
            flush(); i += 1; continue
        if ln.startswith("# "):
            flush(); h = doc.add_heading(ln[2:], level=0 if i == 0 else 1); i += 1; continue
        if ln.startswith("## "):
            flush(); doc.add_heading(ln[3:], level=2); i += 1; continue
        if ln.startswith("### "):
            flush(); doc.add_heading(ln[4:], level=3); i += 1; continue
        if ln.startswith("- "):
            flush(); add_runs(doc.add_paragraph(style="List Bullet"), ln[2:]); i += 1; continue
        if ln[:3] in {f"{n}. " for n in range(1, 10)}:
            flush(); add_runs(doc.add_paragraph(style="List Number"), ln[3:]); i += 1; continue
        if ln.startswith("> "):
            flush()
            t = doc.add_table(rows=1, cols=1); c = t.cell(0, 0); shade(c, "EAF2FB")
            add_runs(c.paragraphs[0], ln[2:]); doc.add_paragraph(); i += 1; continue
        if ln.startswith("| "):
            flush()
            rows = []
            while i < len(lines) and lines[i].startswith("| "):
                rows.append([x.strip() for x in lines[i].strip().strip("|").split(" | ")]); i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    cell = t.cell(r, c); add_runs(cell.paragraphs[0], val, size=10)
                    if r == 0:
                        shade(cell, "0072B2")
                        for run in cell.paragraphs[0].runs:
                            run.bold = True; run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            doc.add_paragraph(); continue
        para_buf.append(ln.strip()); i += 1
    flush()
    dst = os.path.join(OUT, "Team_Briefing_Judge_Prep.docx")
    doc.save(dst)
    with open(os.path.join(OUT, "briefing_part6_technical.html"), "w", encoding="utf-8") as fh:
        fh.write("<html><body>" + to_html(TECH) + "</body></html>")
    words = len(re.findall(r"\w+", TEXT))
    print("saved", dst, "words:", words)


if __name__ == "__main__":
    main()
